import subprocess
import time
import os
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    # 1. Start Django server
    print("Starting Django server...")
    server_process = subprocess.Popen(["python", "manage.py", "runserver"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(5)  # Wait for server to start

    try:
        # 2. Take screenshots using Playwright
        print("Taking screenshots...")
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1280, 'height': 800})
            
            # Login Page
            page.goto("http://localhost:8000/admin/login/")
            time.sleep(2)
            page.screenshot(path="login.png")
            
            # Login
            page.fill("input[name='username']", "pptadmin")
            page.fill("input[name='password']", "pptadminpass")
            page.click("input[type='submit']")
            page.wait_for_load_state("networkidle")
            time.sleep(2)
            
            # Admin Dashboard
            page.screenshot(path="dashboard.png")
            
            # Admission Applications Page
            page.goto("http://localhost:8000/admin/core/admissionapplication/")
            page.wait_for_load_state("networkidle")
            time.sleep(2)
            page.screenshot(path="admissions.png")

            # Hero Section Page
            page.goto("http://localhost:8000/admin/core/herosection/")
            page.wait_for_load_state("networkidle")
            time.sleep(2)
            page.screenshot(path="hero.png")

            browser.close()
            
        print("Screenshots taken successfully.")

        # 3. Create PPTX
        print("Generating PowerPoint Presentation...")
        prs = Presentation()

        # Slide 1: Title
        slide_layout = prs.slide_layouts[0] # Title slide layout
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "College Website - Backend Architecture & Admin Panel"
        subtitle.text = "Powered by Django\nDynamic Content Management"

        # Slide 2: Login
        slide_layout = prs.slide_layouts[5] # Blank slide with title
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Secure Administration Login"
        slide.shapes.add_picture("login.png", Inches(1), Inches(1.5), width=Inches(8))

        # Slide 3: Dashboard
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Admin Dashboard Overview"
        slide.shapes.add_picture("dashboard.png", Inches(1), Inches(1.5), width=Inches(8))

        # Slide 4: Admissions Data
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Managing Admission Applications Data"
        slide.shapes.add_picture("admissions.png", Inches(1), Inches(1.5), width=Inches(8))

        # Slide 5: Hero Section Data
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Managing Dynamic Frontend Content (Hero Section)"
        slide.shapes.add_picture("hero.png", Inches(1), Inches(1.5), width=Inches(8))

        # Slide 6: Backend Tech Stack
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Backend Technologies"
        content = slide.placeholders[1]
        content.text = "Our technology stack ensures security, scalability, and ease of management:\n" \
                       "• Framework: Django (Python)\n" \
                       "• Database: SQLite (Local) / PostgreSQL (Production)\n" \
                       "• File Storage: Cloudinary / Local Media\n" \
                       "• Deployment: Render.com with Gunicorn & WhiteNoise\n" \
                       "• Frontend Integration: Django Templates & Context Processors"

        # Save Presentation
        ppt_filename = "Backend_Presentation.pptx"
        prs.save(ppt_filename)
        print(f"Presentation saved to {ppt_filename}")

    finally:
        # Clean up process
        print("Stopping Django server...")
        server_process.terminate()
        server_process.wait()
        
        # Clean up images
        for img in ["login.png", "dashboard.png", "admissions.png", "hero.png"]:
            if os.path.exists(img):
                os.remove(img)
                
if __name__ == "__main__":
    create_presentation()
